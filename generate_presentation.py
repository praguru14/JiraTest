from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("JiraConf_AI_Overview.pptx")

NAVY = RGBColor(18, 36, 64)
BLUE = RGBColor(35, 103, 178)
TEAL = RGBColor(30, 145, 145)
GOLD = RGBColor(232, 166, 57)
INK = RGBColor(38, 48, 61)
MUTED = RGBColor(92, 105, 121)
PALE = RGBColor(243, 247, 251)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.38, 12.0, 0.55, 27, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 0.98, 11.9, 0.35, 12, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.37), Inches(1.25), Inches(0.07))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()


def add_bullets(slide, items, x, y, w, h, size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(10)
        paragraph.bullet = True
    return box


def add_card(slide, title, body, x, y, w, h, accent=BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = PALE
    card.line.color.rgb = RGBColor(219, 228, 238)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, title, x + 0.25, y + 0.18, w - 0.45, 0.35, 16, NAVY, True)
    add_text(slide, body, x + 0.25, y + 0.62, w - 0.45, h - 0.78, 12, INK)


def add_flow_box(slide, label, x, y, w=1.65, h=0.62, fill=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    add_text(slide, label, x + 0.08, y + 0.18, w - 0.16, 0.25, 12, WHITE, True, PP_ALIGN.CENTER)
    return shape


def add_arrow(slide, x, y, w=0.45):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.32))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = GOLD
    arrow.line.fill.background()


def new_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_title(slide, title, subtitle)
    return slide


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.75), Inches(13.333), Inches(0.75))
    band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background()
    add_text(slide, "JiraConf AI", 0.8, 1.45, 11.8, 0.8, 42, WHITE, True)
    add_text(slide, "AI-assisted release-note automation", 0.84, 2.4, 10.8, 0.55, 24, RGBColor(205, 225, 240))
    add_text(slide, "Project overview, architecture, workflow, and reuse opportunities", 0.86, 3.25, 10.8, 0.4, 16, WHITE)
    add_text(slide, "Jira -> LLM -> Review -> HTML -> Confluence", 0.86, 5.85, 11.4, 0.4, 17, GOLD, True)

    slide = new_slide(prs, "What the project does", "The business outcome")
    add_text(slide, "JiraConf AI turns completed sprint work into publishable customer communication.", 0.8, 1.8, 11.7, 0.55, 24, NAVY, True)
    add_bullets(slide, [
        "Reads completed Jira tickets for a selected sprint.",
        "Uses an LLM to write one customer-friendly note per ticket.",
        "Reviews the generated notes before publication.",
        "Builds an HTML table and creates or updates a Confluence page.",
        "Verifies that the published page contains the expected notes.",
    ], 1.0, 2.65, 11.0, 3.6, 19)

    slide = new_slide(prs, "Architecture", "Adapters, agents, models, and output")
    add_flow_box(slide, "main.py\nor run_agent.py", 0.65, 2.45, 1.8, fill=NAVY)
    add_arrow(slide, 2.55, 2.6)
    add_flow_box(slide, "WorkflowAgent", 3.1, 2.45, 1.8, fill=TEAL)
    add_arrow(slide, 5.0, 2.6)
    add_flow_box(slide, "Planner\nRelease\nReviewer", 5.55, 2.2, 1.8, h=1.12, fill=BLUE)
    add_arrow(slide, 7.45, 2.6)
    add_flow_box(slide, "LLMAdapter", 8.0, 2.45, 1.7, fill=GOLD)
    add_arrow(slide, 9.8, 2.6)
    add_flow_box(slide, "Local / Ollama /\nGemini / Remote", 10.35, 2.2, 2.25, h=1.12, fill=TEAL)
    add_card(slide, "Data adapters", "JiraAgent reads boards, sprints, and Done issues. ConfluenceAgent manages pages and verifies publication.", 1.0, 4.45, 3.45, 1.45, BLUE)
    add_card(slide, "Content services", "PromptLoader, JSON parsing, Jira formatting, HTMLBuilder, and TemplateLoader keep behavior modular.", 4.95, 4.45, 3.45, 1.45, TEAL)
    add_card(slide, "Output", "Confluence storage HTML with a table containing S.No, label, ticket number, and description.", 8.9, 4.45, 3.45, 1.45, GOLD)

    slide = new_slide(prs, "The agent workflow", "A bounded observe-decide-act loop")
    labels = ["FETCH\nTICKETS", "INSPECT\nCONFLUENCE", "GENERATE\nNOTES", "REVIEW\nNOTES", "UPLOAD\nPAGE", "VERIFY\nPAGE", "FINISH"]
    x = 0.55
    for index, label in enumerate(labels):
        add_flow_box(slide, label, x, 2.4, 1.5, 0.78, fill=TEAL if index in (2, 3) else BLUE)
        if index < len(labels) - 1:
            add_arrow(slide, x + 1.57, 2.63, 0.32)
        x += 1.82
    add_text(slide, "The planner proposes the next action. The controller checks prerequisites, limits the run to 12 steps, retries failures, and prevents upload before review.", 1.0, 4.25, 11.3, 0.85, 20, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "This is the project's main agentic behavior.", 2.2, 5.55, 8.9, 0.4, 18, GOLD, True, PP_ALIGN.CENTER)

    slide = new_slide(prs, "How release notes are generated", "The LLM writes content; the template controls layout")
    add_card(slide, "1. Prepare input", "Each Jira issue is formatted with key, type, summary, labels, priority, and description.", 0.8, 1.8, 3.65, 1.55, BLUE)
    add_card(slide, "2. Generate JSON", "release_prompt.txt requires one object: label, ticket_number, and description.", 4.85, 1.8, 3.65, 1.55, TEAL)
    add_card(slide, "3. Validate and retry", "Responses are parsed, checked, and retried. Missing descriptions can fall back to Jira text.", 8.9, 1.8, 3.65, 1.55, GOLD)
    add_card(slide, "4. Review", "The complete array is sent to reviewer_prompt.txt. Valid notes continue to publishing.", 2.8, 4.25, 3.65, 1.55, TEAL)
    add_card(slide, "5. Render", "HTMLBuilder loads release_notes.html and replaces {{SPRINT_NAME}} and {{TABLE_ROWS}}.", 6.85, 4.25, 3.65, 1.55, BLUE)

    slide = new_slide(prs, "The agents", "What each component is responsible for")
    add_card(slide, "WorkflowAgent", "Orchestrates actions, state, recovery, safety checks, and verification.", 0.7, 1.7, 3.8, 1.35, TEAL)
    add_card(slide, "PlannerAgent", "Asks the LLM to choose the next available workflow action.", 4.75, 1.7, 3.8, 1.35, BLUE)
    add_card(slide, "ReleaseAgent", "Creates one structured note per Jira issue, with retries and validation.", 8.8, 1.7, 3.8, 1.35, GOLD)
    add_card(slide, "ReviewerAgent", "Checks the full note set against the release-note policy.", 0.7, 4.05, 3.8, 1.35, GOLD)
    add_card(slide, "JiraAgent", "Wraps the Jira Python client and exposes project-specific data operations.", 4.75, 4.05, 3.8, 1.35, TEAL)
    add_card(slide, "ConfluenceAgent", "Wraps the Confluence client and handles page creation, update, and checks.", 8.8, 4.05, 3.8, 1.35, BLUE)

    slide = new_slide(prs, "Models and prompts", "The model is replaceable")
    add_bullets(slide, [
        "LLMAdapter selects the provider using LLM_PROVIDER.",
        "Local: Hugging Face Transformers and PyTorch.",
        "Ollama: local HTTP endpoint, configurable model.",
        "Gemini: Google REST API using GEMINI_API_KEY.",
        "Remote: Hugging Face-style endpoint using HF_API_URL and HF_API_TOKEN.",
        "Planner, release, and reviewer behavior is controlled by text files in prompts/.",
    ], 0.9, 1.75, 7.1, 4.8, 18)
    add_card(slide, "Key design idea", "The rest of the application uses one generate(system_prompt, user_prompt) interface. Changing the model provider does not require rewriting the agents.", 8.55, 2.15, 3.75, 2.2, TEAL)

    slide = new_slide(prs, "What MCP would change", "MCP is a tool connection standard")
    add_flow_box(slide, "WorkflowAgent", 0.9, 2.45, 2.1, fill=NAVY)
    add_arrow(slide, 3.25, 2.68, 0.55)
    add_flow_box(slide, "MCP client", 4.0, 2.45, 2.1, fill=TEAL)
    add_arrow(slide, 6.35, 2.68, 0.55)
    add_flow_box(slide, "Jira MCP\nConfluence MCP", 7.1, 2.25, 2.3, h=1.18, fill=BLUE)
    add_arrow(slide, 9.65, 2.68, 0.55)
    add_flow_box(slide, "External\nservices", 10.4, 2.25, 2.0, h=1.18, fill=GOLD)
    add_text(slide, "MCP would standardize how tools are discovered and called. It would not automatically make the project more intelligent or autonomous.", 1.0, 4.3, 11.3, 0.9, 21, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "Current project: direct Jira and Confluence API clients.\nPossible future: replace those adapters with MCP clients.", 2.0, 5.55, 9.3, 0.65, 16, MUTED, False, PP_ALIGN.CENTER)

    slide = new_slide(prs, "Reuse opportunities", "The same pattern can support other workflows")
    add_bullets(slide, [
        "Incident summaries and postmortems",
        "Customer-support response generation",
        "Sprint and project status reports",
        "Product documentation and changelog generation",
        "Ticket classification and prioritization",
        "Compliance, audit, or security reports",
        "Publishing to email, Slack, SharePoint, databases, or other APIs",
    ], 1.0, 1.75, 6.4, 4.8, 19)
    add_card(slide, "What to customize", "Replace the data adapter, change the prompts and JSON schema, update the output template, and redefine the workflow actions.", 8.2, 2.35, 3.8, 2.25, GOLD)

    slide = new_slide(prs, "Current limitations", "Important before production reuse")
    add_bullets(slide, [
        "Jira query is hard-coded to a project, sprint, Done status, and 100 issues.",
        "Existing-page checks compare row counts, not ticket identities.",
        "HTML values are inserted without escaping.",
        "Updating an existing page requires interactive confirmation.",
        "Testing is limited; test_llm.py is a smoke script rather than a test suite.",
        "The architecture diagram is stale and the README needs repair.",
    ], 0.9, 1.7, 7.4, 4.9, 18)
    add_card(slide, "Recommended next steps", "Add structured tests, escape HTML, make integrations configurable, support headless updates, and compare ticket IDs during verification.", 8.7, 2.35, 3.45, 2.35, TEAL)

    slide = new_slide(prs, "Summary", "The reusable idea")
    add_text(slide, "JiraConf AI is more than a release-note template.", 0.9, 1.8, 11.5, 0.55, 27, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "It is a planner-driven workflow engine with replaceable data sources, LLM providers, prompts, validators, and output templates.", 1.3, 2.75, 10.7, 0.95, 23, INK, False, PP_ALIGN.CENTER)
    add_text(slide, "MCP can make the tool connections more portable.\nThe agentic behavior comes from the observe -> decide -> act -> verify loop.", 1.8, 4.35, 9.8, 0.8, 20, TEAL, True, PP_ALIGN.CENTER)
    add_text(slide, "Source: JiraConf project files and PROJECT_GUIDE.md", 0.8, 6.65, 11.7, 0.3, 11, MUTED, False, PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()